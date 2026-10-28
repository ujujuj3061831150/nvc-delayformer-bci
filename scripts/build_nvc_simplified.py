from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pathlib import Path
P=Presentation();P.slide_width=Inches(16.38);P.slide_height=Inches(5.98);sl=P.slides.add_slide(P.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=RGBColor.from_string('FFFFFF')
def line(x1,y1,x2,y2,arrow=False):
 s=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));s.line.color.rgb=RGBColor.from_string('3B3B3B');s.line.width=Pt(1.4)
 if arrow:
  e=OxmlElement('a:tailEnd');e.set('type','triangle');s._element.spPr.ln.append(e)
 return s
def text(t,x,y,w,h,pt=16,bold=False):
 s=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=s.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=tf.margin_right=Pt(1);tf.margin_top=tf.margin_bottom=Pt(0);tf.vertical_anchor=MSO_ANCHOR.MIDDLE;p=tf.paragraphs[0];p.text=t;p.alignment=PP_ALIGN.CENTER;p.font.name='Times New Roman';p.font.size=Pt(pt);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string('171717');return s
def box(x,y,w,h,fill='F8F8F8'):
 s=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));s.adjustments[0]=.06;s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string(fill);s.line.color.rgb=RGBColor.from_string('454545');s.line.width=Pt(1.1)
 for e in s._element.xpath('.//a:effectRef'):e.set('idx','0')
 return s
# Connections are drawn before boxes so nodes cover crossing line ends.
line(.90,1.82,1.35,1.82,True);line(.90,3.18,1.35,3.18,True)
line(3.18,1.82,4.22,1.82);line(4.22,1.82,4.22,2.12);line(4.22,2.12,4.48,2.12,True)
line(3.18,3.18,4.22,3.18);line(4.22,3.18,4.22,2.73);line(4.22,2.73,4.48,2.73,True)
line(5.53,1.14,5.53,1.85,True)
line(6.58,2.43,7.32,2.43,True)
# EEG encoded features also enter deep pooling.
line(3.65,1.82,3.65,1.42);line(3.65,1.42,8.20,1.42);line(8.20,1.42,8.20,2.10,True)
# Original input signals feed shallow direct features.
line(1.08,1.82,1.08,4.65);line(1.08,4.65,4.48,4.65,True);line(1.08,3.18,1.08,4.65)
line(6.58,4.65,7.32,4.65,True)
# Both branch embeddings enter entropy-guided fusion.
line(9.12,2.43,10.08,2.43);line(10.08,2.43,10.08,3.12);line(10.08,3.12,10.48,3.12,True)
line(9.12,4.65,10.08,4.65);line(10.08,4.65,10.08,3.69);line(10.08,3.69,10.48,3.69,True)
line(12.75,3.40,13.45,3.40,True)
# Input, branch modules and outputs.
text('EEG',.16,1.60,.72,.42,18,True);text('fNIRS',.05,2.97,.85,.42,18,True)
box(1.35,1.46,1.83,.72,'F4F4F4');text('EEG encoder',1.40,1.58,1.73,.46,17,True)
box(1.35,2.82,1.83,.72,'F4F4F4');text('fNIRS encoder',1.40,2.94,1.73,.46,17,True)
box(4.48,.45,2.10,.69,'FFFFFF');text('Delay prior + mask',4.54,.56,1.98,.45,16,True)
box(4.48,1.85,2.10,1.15,'F0F0F0');text('Cross-attention',4.55,2.18,1.96,.48,17,True)
text('K, V',3.70,1.95,.50,.30,14);text('Q',3.72,2.84,.36,.30,14)
box(7.32,2.10,1.80,.67,'F4F4F4');text('Temporal pooling',7.37,2.21,1.70,.43,16,True)
box(4.48,4.28,2.10,.74,'F5F5F5');text('Spectral + optical\nfeatures',4.52,4.37,2.02,.56,16,True)
box(7.32,4.28,1.80,.74,'F4F4F4');text('Projection',7.37,4.44,1.70,.40,17,True)
box(10.48,2.82,2.27,1.17,'EAEAEA');text('Entropy-guided\ndual-stream fusion',10.56,3.02,2.11,.75,17,True)
box(13.45,3.04,1.77,.72,'F7F7F7');text('Classifier',13.50,3.17,1.67,.46,17,True)
text('Deep stream',7.56,2.89,1.35,.35,14);text('Shallow stream',7.45,5.10,1.65,.35,14)
for sh in sl.shapes:
 for e in sh._element.xpath('.//a:effectRef'): e.set('idx','0')
Path('reports').mkdir(exist_ok=True);P.save('reports/NVC_architecture_simplified.pptx');print('Saved')

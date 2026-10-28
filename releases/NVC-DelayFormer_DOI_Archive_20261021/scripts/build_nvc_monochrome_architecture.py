from pathlib import Path
from xml.sax.saxutils import escape
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
ROOT=Path(__file__).resolve().parents[1]; out=ROOT/'reports';out.mkdir(exist_ok=True)
prs=Presentation();prs.slide_width=Inches(14);prs.slide_height=Inches(5.25);slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor(255,255,255)
svg=['<svg xmlns="http://www.w3.org/2000/svg" width="1400" height="525" viewBox="0 0 1400 525">','<title>NVC-DelayFormer architecture</title>','<desc>EEG and fNIRS encoders, delay-prior cross-attention, parallel direct features and entropy-guided fusion.</desc>','<rect width="1400" height="525" fill="white"/>','<defs><marker id="a" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="5" markerHeight="5" orient="auto-start-reverse"><path d="M 0 0 L 10 5 L 0 10 z" fill="#353535"/></marker></defs>']
def rect(x,y,w,h,fill='#fafafa',stroke='#454545',radius=4):
 sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x/100), Inches(y/100), Inches(w/100), Inches(h/100));sh.adjustments[0]=.05;sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill[1:]);sh.line.color.rgb=RGBColor.from_string(stroke[1:]);sh.line.width=Pt(.9)
 svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="{radius}" fill="{fill}" stroke="{stroke}" stroke-width="1.3"/>')
def text(x,y,w,h,content,pt=12,bold=False,color='#272727',align='center'):
 sh=slide.shapes.add_textbox(Inches(x/100),Inches(y/100),Inches(w/100),Inches(h/100)); tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=Pt(.5);tf.margin_right=Pt(.5);tf.margin_top=Pt(0);tf.margin_bottom=Pt(0);tf.vertical_anchor=MSO_ANCHOR.MIDDLE
 p=tf.paragraphs[0];p.text=content;p.font.name='Times New Roman';p.font.size=Pt(pt);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color[1:]);p.alignment=PP_ALIGN.CENTER if align=='center' else PP_ALIGN.LEFT
 rows=content.split('\n');fs=pt*100/72; spacing=fs*1.12; y0=y+h/2-(len(rows)-1)*spacing/2+fs*.34;tx=x+w/2 if align=='center' else x+1; anchor='middle' if align=='center' else 'start'
 svg.append(f'<text x="{tx}" y="{y0:.2f}" font-family="Times New Roman,serif" font-size="{fs:.2f}" font-weight="{700 if bold else 400}" fill="{color}" text-anchor="{anchor}">'+''.join(f'<tspan x="{tx}" dy="{0 if i==0 else spacing:.2f}">{escape(row)}</tspan>' for i,row in enumerate(rows))+'</text>')
def arrow(points,end=True,dashed=False):
 for a,b in zip(points,points[1:]):
  sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(a[0]/100),Inches(a[1]/100),Inches(b[0]/100),Inches(b[1]/100));sh.line.color.rgb=RGBColor.from_string('353535');sh.line.width=Pt(1.0)
  if end and b==points[-1]:
   e=OxmlElement('a:tailEnd');e.set('type','triangle');sh._element.spPr.ln.append(e)
 pts=' '.join(f'{x},{y}' for x,y in points);svg.append(f'<polyline points="{pts}" fill="none" stroke="#353535" stroke-width="1.5"'+(' marker-end="url(#a)"' if end else '')+(' stroke-dasharray="4 4"' if dashed else '')+'/>')
# Horizontal, monochrome layout adapted from the supplied source.
text(24,15,290,28,'Multimodal inputs',13,True);text(380,15,350,28,'Physiology-guided temporal fusion',13,True);text(850,15,225,28,'Dual-stream fusion',13,True);text(1120,15,240,28,'Classification',13,True)
# input + separate EA / CrEA. EEG (not fNIRS) is aligned.
text(8,85,88,42,'EEG\n30 channels',12,True);rect(104,84,110,48);text(108,89,102,38,'EA / CrEA',12,True);arrow([(94,108),(103,108)])
rect(255,80,130,56,'#eeeeee');text(259,83,122,50,'2D-EEGNet\nencoder',12,True);arrow([(214,108),(254,108)])
text(12,227,80,52,'fNIRS\n72 features',12,True);rect(255,224,130,56,'#eeeeee');text(259,227,122,50,'Conv1D\nencoder',12,True);arrow([(94,251),(254,251)])
text(386,77,73,28,'Hₑ → K, V',11);text(386,263,73,27,'Hₙ → Q',11)
# cross attention, TPP receives EEG features and enhanced fNIRS features.
rect(489,138,196,113,'#f3f3f3');text(495,144,184,35,'Cross-attention',13,True);text(495,182,184,31,'Softmax(QKᵀ/√d + P + M)',11.2);text(497,216,180,22,'4 heads',10.5)
arrow([(385,108),(451,108),(451,164),(488,164)]);arrow([(385,251),(488,224)])
rect(492,281,91,48,'#ffffff');text(496,286,83,38,'Delay prior\nP',11,True);rect(591,281,91,48,'#ffffff');text(595,286,83,38,'Timestamp\nmask M',11,True)
arrow([(537,280),(537,252)]);arrow([(636,280),(636,252)]);text(491,339,195,25,'τₕ = 3.5, 4.5, 5.5, 6.5 s',10.5)
rect(737,143,146,102,'#eeeeee');text(741,150,138,40,'Attentive\ntemporal pooling',12,True);text(741,198,138,33,'Deep projection\n128 → 64',10.5);arrow([(685,194),(736,194)])
arrow([(425,108),(425,63),(810,63),(810,142)]);text(560,65,174,22,'EEG tokens',10.5)
text(762,251,100,24,'z_deep',11,True)
# direct stream operates on aligned EEG and optical inputs, not learned band encoder
rect(455,402,231,74,'#f6f6f6');text(459,406,223,65,'Direct features\nEEG θ / α / β power\nfNIRS mean and slope',11,True)
arrow([(232,108),(232,380),(478,380),(478,401)]);arrow([(179,251),(179,394),(447,394),(447,438),(454,438)])
rect(737,408,146,62,'#eeeeee');text(741,412,138,52,'Shallow projection\n234 → 32 → 64',11,True);arrow([(686,439),(736,439)]);text(761,472,100,23,'z_shallow',11,True)
# entropy-gate and final classification. Two arrows converge on gate.
rect(944,225,170,104,'#dedede');text(949,232,160,34,'Entropy-guided gate',12,True);text(950,268,158,53,'feature-wise sigmoid\ndeep + shallow',11)
arrow([(883,194),(915,194),(915,252),(943,252)]);arrow([(883,439),(915,439),(915,305),(943,305)])
rect(1190,243,135,69,'#eeeeee');text(1194,247,127,61,'Linear classifier\n2 classes',12,True);arrow([(1114,277),(1189,277)])
text(110,146,107,25,'trial-wise or offline',9.8);text(1139,312,206,36,'joint logits',10.5)
# avoid theme effects in PowerPoint
for sh in slide.shapes:
 for el in sh._element.xpath('.//a:effectRef'):el.set('idx','0')
svg.append('</svg>'); (out/'NVC_DelayFormer_architecture_monochrome.svg').write_text('\n'.join(svg),encoding='utf-8');prs.save(out/'NVC_DelayFormer_architecture_monochrome.pptx')
print('Created',len(slide.shapes),'editable shapes')

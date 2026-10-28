from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
P=Presentation('reports/architecture_original_text_simplified_clean.pptx');s=P.slides[0]
# Retain the original encoders, cross-attention illustration, enclosing frame and typography.
for sh in list(s.shapes):
 if sh.left/Inches(1)>=7.85: sh._element.getparent().remove(sh._element)
def line(x1,y1,x2,y2,arrow=False):
 z=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));z.line.color.rgb=RGBColor.from_string('252525');z.line.width=Pt(1.35)
 if arrow: e=OxmlElement('a:tailEnd');e.set('type','triangle');z._element.spPr.ln.append(e)
 for q in z._element.xpath('.//a:effectRef'):q.set('idx','0')
 return z
def txt(t,x,y,w,h,sz=15,bold=False):
 z=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));f=z.text_frame;f.clear();f.word_wrap=True;f.margin_left=f.margin_right=Pt(1);f.margin_top=f.margin_bottom=Pt(0);f.vertical_anchor=MSO_ANCHOR.MIDDLE;p=f.paragraphs[0];p.text=t;p.alignment=PP_ALIGN.CENTER;p.font.name='Times New Roman';p.font.size=Pt(sz);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string('171717');return z
def box(x,y,w,h,fill='F7F7F7'):
 z=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));z.adjustments[0]=.07;z.fill.solid();z.fill.fore_color.rgb=RGBColor.from_string(fill);z.line.color.rgb=RGBColor.from_string('363636');z.line.width=Pt(1.2)
 for q in z._element.xpath('.//a:effectRef'):q.set('idx','0')
 return z
# Arrows express the two independent representations and gate-controlled fusion.
line(9.55,1.69,10.00,1.69,True)
line(10.90,1.69,10.90,2.57,True)
line(9.55,4.35,10.00,4.35,True)
line(10.90,4.03,10.90,3.72,True)
line(11.80,3.15,12.24,3.15,True)
line(9.55,1.69,12.05,1.69);line(12.05,1.69,12.05,2.91);line(12.05,2.91,12.24,2.91,True)
line(11.80,4.35,12.05,4.35);line(12.05,4.35,12.05,3.39);line(12.05,3.39,12.24,3.39,True)
line(13.82,3.15,14.05,3.15,True);line(15.29,3.15,15.65,3.15,True)
line(4.41,1.48,4.41,0.98);line(4.41,0.98,8.79,0.98);line(8.79,0.98,8.79,1.35,True)
# Explicit EEG/fNIRS source is shown as a label; no false assertion of direct cross-attention input.
box(8.04,1.35,1.51,.70);txt('Temporal\npooling',8.08,1.43,1.43,.53,16,True)
box(8.04,3.99,1.51,.72);txt('EEG + fNIRS\nfeatures',8.08,4.06,1.43,.58,15,True)
box(10.00,3.99,1.80,.72);txt('Shallow\nprojection',10.05,4.06,1.70,.58,15,True)
box(10.00,2.57,1.80,1.15,'EDEDED');txt('Entropy-guided\ngate',10.05,2.76,1.70,.74,16,True)
box(12.24,2.70,1.58,.90,'EDEDED');txt('Feature-wise\nfusion',12.30,2.83,1.46,.60,15,True)
box(14.05,2.79,1.24,.72);txt('Classifier',14.08,2.94,1.18,.42,15,True)
txt('Deep stream',8.12,2.16,1.34,.27,12);txt('Spectral + optical',8.00,4.82,1.60,.30,12)
P.save('reports/architecture_original_flow_corrected.pptx');print('saved')
